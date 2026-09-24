"""
Competitor Price Tracker
========================
Scrapes every product from an online catalogue, compares the prices with the
previous run, and produces a formatted Excel report.

Practice target: https://books.toscrape.com (a site built for scraping practice).
To adapt it for a real client, change the selectors inside `parse_books()`.

Usage:
    python price_tracker.py                 # scrape all pages
    python price_tracker.py --pages 5       # only the first 5 pages (quick test)
    python price_tracker.py --demo          # pretend some prices changed (for demos)
"""

import argparse
import json
import logging
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

# --------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------
DEFAULT_BASE_URL = "https://books.toscrape.com/"
HISTORY_FILE = Path("data/last_run.json")
DEFAULT_OUTPUT = Path("output/price_report.xlsx")
HEADERS = {"User-Agent": "Mozilla/5.0 (portfolio project: price tracker)"}
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

# Excel styling
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
COLUMN_WIDTHS = {
    "Title": 55, "Price (£)": 12, "Rating": 9, "Stock": 12, "URL": 45,
    "Scraped At": 18, "Old Price (£)": 14, "New Price (£)": 14,
    "Change (£)": 12, "Change (%)": 12,
}

log = logging.getLogger("price_tracker")


# --------------------------------------------------------------------------
# Step 1: Download pages
# --------------------------------------------------------------------------
def fetch(session, url, retries=3, timeout=15):
    """Download one page. Returns the HTML text, or None if the page does not exist (404)."""
    for attempt in range(1, retries + 1):
        try:
            response = session.get(url, timeout=timeout)
            if response.status_code == 404:
                return None
            response.raise_for_status()
            response.encoding = "utf-8"  # otherwise the £ sign can turn into "Â£"
            return response.text
        except requests.RequestException as error:
            log.warning("Attempt %d/%d failed for %s (%s)", attempt, retries, url, error)
            time.sleep(2 * attempt)
    raise RuntimeError(f"Could not download {url} after {retries} attempts")


# --------------------------------------------------------------------------
# Step 2: Pull the data out of the HTML
# --------------------------------------------------------------------------
def parse_price(text):
    """'£51.77' -> 51.77"""
    match = re.search(r"\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def parse_books(html, page_url):
    """Return a list of dicts, one per book found on the page."""
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for item in soup.select("article.product_pod"):
        link = item.h3.a
        books.append({
            "Title": link["title"],
            "Price (£)": parse_price(item.select_one(".price_color").text),
            "Rating": RATING_WORDS.get(item.select_one("p.star-rating")["class"][1], 0),
            "Stock": item.select_one(".availability").get_text(strip=True),
            "URL": urljoin(page_url, link["href"]),
        })
    return books


def scrape_all(base_url, max_pages, delay):
    """Visit page after page until the site says 'no more pages'."""
    session = requests.Session()
    session.headers.update(HEADERS)
    all_books = []
    for page in range(1, max_pages + 1):
        url = f"{base_url}catalogue/page-{page}.html"
        html = fetch(session, url)
        if html is None:
            log.info("Page %d does not exist - reached the last page.", page)
            break
        books = parse_books(html, url)
        all_books.extend(books)
        log.info("Page %d: %d books (total so far: %d)", page, len(books), len(all_books))
        time.sleep(delay)  # be polite: do not hammer the server
    return all_books


# --------------------------------------------------------------------------
# Step 3: Compare with the previous run
# --------------------------------------------------------------------------
def load_history():
    if HISTORY_FILE.exists():
        return json.loads(HISTORY_FILE.read_text(encoding="utf-8")).get("prices", {})
    return {}


def save_history(df, scraped_at):
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    data = {"scraped_at": scraped_at, "prices": dict(zip(df["URL"], df["Price (£)"]))}
    HISTORY_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")


def simulate_previous_prices(df, share=0.08, seed=42):
    """DEMO ONLY: invent 'yesterday's prices' so the change detection can be shown."""
    rng = random.Random(seed)
    old = dict(zip(df["URL"], df["Price (£)"]))
    for url in rng.sample(list(old), k=max(1, int(len(old) * share))):
        old[url] = round(old[url] * rng.choice([0.8, 0.9, 1.15, 1.25]), 2)
    return old


def find_price_changes(df, old_prices):
    columns = ["Title", "Old Price (£)", "New Price (£)", "Change (£)", "Change (%)", "URL"]
    rows = []
    for _, book in df.iterrows():
        old = old_prices.get(book["URL"])
        new = book["Price (£)"]
        if old is not None and abs(new - old) > 0.001:
            rows.append({
                "Title": book["Title"],
                "Old Price (£)": old,
                "New Price (£)": new,
                "Change (£)": round(new - old, 2),
                "Change (%)": round((new - old) / old * 100, 1),
                "URL": book["URL"],
            })
    return pd.DataFrame(rows, columns=columns)


# --------------------------------------------------------------------------
# Step 4: Build the Excel report
# --------------------------------------------------------------------------
def style_table(ws):
    """Header colours, fonts, column widths, number formats, clickable links, filters."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    headers = {cell.value: cell.column_letter for cell in ws[1]}

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT

    for name, letter in headers.items():
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, 16)
        if "£" in name:
            for cell in ws[letter][1:]:
                cell.number_format = "£#,##0.00"
        if name == "Change (%)":
            for cell in ws[letter][1:]:
                cell.number_format = '0.0"%"'
        if name == "URL":
            for cell in ws[letter][1:]:
                cell.hyperlink = cell.value
                cell.font = LINK_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return headers


def colour_price_changes(ws, headers):
    """Red = price went up, green = price went down."""
    if ws.max_row < 2:
        return
    letter = headers["Change (£)"]
    cells = f"{letter}2:{letter}{ws.max_row}"
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ws.conditional_formatting.add(cells, CellIsRule(operator="greaterThan", formula=["0"], fill=red))
    ws.conditional_formatting.add(cells, CellIsRule(operator="lessThan", formula=["0"], fill=green))


def build_summary(wb, n_books, n_changes, n_deals, meta):
    """Summary sheet. Numbers are Excel formulas, so they update if the data is edited."""
    ws = wb.create_sheet("Summary", 0)
    last = max(n_books + 1, 2)
    books = "'All Books'"
    rows = [
        ("Total books", f"=COUNTA({books}!A2:A{last})"),
        ("Average price (£)", f"=AVERAGE({books}!B2:B{last})"),
        ("Median price (£)", f"=MEDIAN({books}!B2:B{last})"),
        ("Cheapest price (£)", f"=MIN({books}!B2:B{last})"),
        ("Most expensive price (£)", f"=MAX({books}!B2:B{last})"),
        ("Books rated 4-5 stars", f'=COUNTIF({books}!C2:C{last},">=4")'),
        ("Books in stock", f'=COUNTIF({books}!D2:D{last},"In stock")'),
        ("Price changes detected", f"=COUNTA('Price Changes'!A2:A{max(n_changes + 1, 2)})"),
        ("Top deals found", f"=COUNTA('Top Deals'!A2:A{max(n_deals + 1, 2)})"),
        ("Scraped at", meta["scraped_at"]),
        ("Source", meta["source"]),
        ("Data mode", meta["mode"]),
    ]
    ws.append(["Metric", "Value"])
    for row in rows:
        ws.append(list(row))

    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            cell.font = BODY_FONT
        row[1].alignment = Alignment(horizontal="left")
    for row_number in (3, 4, 5, 6):  # the price rows
        ws.cell(row=row_number, column=2).number_format = "£#,##0.00"

    # Rating breakdown table + bar chart
    table_top = len(rows) + 4
    ws.cell(row=table_top, column=1, value="Rating")
    ws.cell(row=table_top, column=2, value="Books")
    for cell in (ws.cell(row=table_top, column=1), ws.cell(row=table_top, column=2)):
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for stars in range(1, 6):
        r = table_top + stars
        ws.cell(row=r, column=1, value=f"{stars} star" + ("s" if stars > 1 else "")).font = BODY_FONT
        cell = ws.cell(row=r, column=2, value=f"=COUNTIF({books}!C2:C{last},{stars})")
        cell.font = BODY_FONT
        cell.alignment = Alignment(horizontal="left")

    chart = BarChart()
    chart.type = "col"
    chart.title = "Books by rating"
    chart.legend = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(ws, min_col=2, min_row=table_top, max_row=table_top + 5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=table_top + 1, max_row=table_top + 5))
    chart.height, chart.width = 8, 14
    ws.add_chart(chart, "D2")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42


def write_excel(path, df, changes, deals, meta):
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Books", index=False)
        changes.to_excel(writer, sheet_name="Price Changes", index=False)
        deals.to_excel(writer, sheet_name="Top Deals", index=False)
        wb = writer.book
        style_table(wb["All Books"])
        colour_price_changes(wb["Price Changes"], style_table(wb["Price Changes"]))
        style_table(wb["Top Deals"])
        build_summary(wb, len(df), len(changes), len(deals), meta)


# --------------------------------------------------------------------------
# Main program
# --------------------------------------------------------------------------
def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Scrape prices and build an Excel report.")
    parser.add_argument("--pages", type=int, default=50, help="maximum pages to scrape (default: 50)")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Excel file to create")
    parser.add_argument("--deal-price", type=float, default=20.0, help="max price for 'Top Deals' (default: 20)")
    parser.add_argument("--delay", type=float, default=0.3, help="seconds to wait between pages")
    parser.add_argument("--demo", action="store_true", help="simulate previous prices to show change detection")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help="site to scrape")
    return parser.parse_args(argv)


def main(argv=None):
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(message)s", datefmt="%H:%M:%S")
    args = parse_args(argv)
    base_url = args.base_url.rstrip("/") + "/"

    log.info("Starting scrape of %s", base_url)
    books = scrape_all(base_url, args.pages, args.delay)
    if not books:
        log.error("No books found. Check the URL or the selectors in parse_books().")
        return 1

    scraped_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    df = pd.DataFrame(books)
    df["Scraped At"] = scraped_at

    if args.demo:
        log.warning("DEMO MODE: previous prices are simulated, not real history.")
        old_prices = simulate_previous_prices(df)
    else:
        old_prices = load_history()
        if not old_prices:
            log.info("First run: no previous prices to compare with yet.")

    changes = find_price_changes(df, old_prices)
    deals = (
        df[(df["Price (£)"] <= args.deal_price) & (df["Rating"] >= 4)]
        .drop(columns="Scraped At")
        .sort_values(["Price (£)", "Title"])
    )

    meta = {
        "scraped_at": scraped_at,
        "source": base_url,
        "mode": "DEMO - previous prices simulated" if args.demo else "Live data",
    }
    write_excel(args.output, df, changes, deals, meta)
    save_history(df, scraped_at)

    log.info("Done! %d books, %d price changes, %d top deals.", len(df), len(changes), len(deals))
    log.info("Report saved to: %s", args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
